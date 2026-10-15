"""DQ excel 读取器测试（v2 波 1b 收尾：dq模板.xlsx 落地）。

覆盖：
1. read_dq_excel 同构产出（与 yml 解析字段一致 + target_schema 特有）
2. 容错（列模糊匹配/无编号行跳过/无主 sheet/文件不存在）
3. main --dq 通道 e2e
4. manifest provided 解锁（含空 excel fail loud）

运行:
    pytest tests/test_dq_excel.py -v
"""

import json
import shutil
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ANALYZER_REF = PROJECT_ROOT / "dws-pipeline-analyzer" / "references"
DOCS = PROJECT_ROOT / "docs"
sys.path.insert(0, str(ANALYZER_REF))

from inputs import read_dq_excel

ET_XLSX = DOCS / "术加制品参考案例" / "execution_tasks_DWB_CUST_ASSET_D.xlsx"


def _write_dq_excel(path, rows):
    """构造 DQ 平台导出格式 excel（主 sheet「规则基础信息」+ 附属空 sheet）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "规则基础信息"
    ws.append(["规则Number", "规则ID", "操作类型", "规则中文名", "规则中文描述",
               "检查类型", "规则负责人", "目标表schema", "目标表", "目标表编码",
               "SQL", "告警级别"])
    for r in rows:
        ws.append(r)
    wb.create_sheet("参数配置").append(["ID", "规则ID", "参数名"])
    wb.create_sheet("输出列配置").append(["ID", "规则ID", "列原名"])
    wb.save(str(path))


class TestReadDqExcel:
    def test_isomorphic_fields_with_schema(self, tmp_path):
        xlsx = tmp_path / "dq.xlsx"
        _write_dq_excel(xlsx, [
            ["R001", "1", "新增", "空值检查", "主键非空校验", "空值检查",
             "zhangsan", "dws", "dwb_cust_asset_d", "T001",
             "SELECT COUNT(*) FROM dws.dwb_cust_asset_d WHERE cust_id IS NULL", "严重"],
        ])
        rules = read_dq_excel(str(xlsx))
        assert rules is not None and len(rules) == 1
        r = rules[0]
        # 与 yml 解析产物同构（loader 接口契约）+ excel 特有 target_schema
        assert set(r.keys()) == {
            "rule_number", "rule_name", "rule_desc", "check_type", "alert_level",
            "sql", "target_table", "target_schema", "owner", "dq_file",
        }
        assert r["rule_number"] == "R001"
        assert r["rule_name"] == "空值检查"
        assert r["target_schema"] == "dws"
        assert r["target_table"] == "dwb_cust_asset_d"
        assert "cust_id IS NULL" in r["sql"]
        assert r["alert_level"] == "严重"
        assert r["dq_file"] == str(xlsx)

    def test_multi_rows_and_skip_invalid(self, tmp_path):
        """多行规则全收；无规则Number的行（模板尾部统计行）跳过。"""
        xlsx = tmp_path / "dq.xlsx"
        _write_dq_excel(xlsx, [
            ["R001", "1", "新增", "规则一", "描述一", "空值检查", "u1", "dws", "t1", "T1", "SELECT 1", "提示"],
            [None, None, None, None, None, None, None, None, None, None, None, None],
            ["R002", "2", "新增", "规则二", "描述二", "唯一性", "u2", "dws", "t1", "T2", "SELECT 2", "严重"],
        ])
        rules = read_dq_excel(str(xlsx))
        assert [r["rule_number"] for r in rules] == ["R001", "R002"]

    def test_fuzzy_header_matching(self, tmp_path):
        """列名全角/空格变体容错（平台导出工具格式不统一）。"""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "规则基础信息"
        ws.append(["规则 Number", "规则中文名", "规则中文描述", "检查类型",
                   "目标表schema", "目标表", "SQL", "告警级别", "规则负责人"])
        ws.append(["R1", "名", "述", "型", "dws", "t1", "SELECT 1", "级", "人"])
        xlsx = tmp_path / "dq2.xlsx"
        wb.save(str(xlsx))
        rules = read_dq_excel(str(xlsx))
        assert rules[0]["rule_number"] == "R1"
        assert rules[0]["target_table"] == "t1"

    def test_invalid_inputs(self, tmp_path):
        assert read_dq_excel(tmp_path / "nope.xlsx") is None
        import openpyxl
        xlsx = tmp_path / "bad.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "别的sheet"
        wb.save(str(xlsx))
        assert read_dq_excel(str(xlsx)) is None

    def test_header_only_returns_none(self, tmp_path):
        """只有表头无数据行（= 当前仓内模板形态）→ None（给了但读不到规则）。"""
        xlsx = tmp_path / "empty.xlsx"
        _write_dq_excel(xlsx, [])
        assert read_dq_excel(str(xlsx)) is None


def _run_main(argv):
    from analyzer import main
    import io, contextlib
    old_argv = sys.argv
    sys.argv = ["analyzer.py"] + argv
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            main()
        return buf.getvalue()
    finally:
        sys.argv = old_argv


def _run_facts(argv):
    from facts import main
    import io, contextlib
    old_argv = sys.argv
    sys.argv = ["facts.py"] + argv
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            main()
        return buf.getvalue()
    finally:
        sys.argv = old_argv


class TestDqChannelsE2E:
    def test_main_dq_flag(self, tmp_path):
        """/--dq 通道：xlsx 制品 + DQ excel → meta.dq_rules 注入 + 契约投影。"""
        from baseline_export import export_baseline_v1
        from inputs import build_baseline_export_ctx
        xlsx = tmp_path / "dq.xlsx"
        _write_dq_excel(xlsx, [
            ["R001", "1", "新增", "空值检查", "主键非空", "空值检查",
             "u", "dws", "dwb_cust_asset_d", "T1", "SELECT 1", "严重"],
        ])
        _run_main(["--input", str(ET_XLSX), "--output", str(tmp_path / "out"),
                   "--dq", str(xlsx)])
        draft = json.loads((tmp_path / "out" / "DWB_CUST_ASSET_D" / "knowledge_draft.json")
                           .read_text(encoding="utf-8"))
        assert len(draft["meta"]["dq_rules"]) == 1
        assert draft["meta"]["dq_rules"][0]["target_schema"] == "dws"

    def test_facts_dir_dq_provided(self, tmp_path):
        """facts 目录直入：DQ 制品分拣 → 契约 dq_rules 含 target_schema。"""
        bag = tmp_path / "bag"
        bag.mkdir()
        shutil.copy(ET_XLSX, bag / "et.xlsx")
        _write_dq_excel(bag / "dq.xlsx", [
            ["R001", "1", "新增", "空值检查", "主键非空", "空值检查",
             "u", "dws", "dwb_cust_asset_d", "T1", "SELECT 1", "严重"],
        ])
        out = tmp_path / "o" / "asset_facts.json"
        _run_facts(["--dir", str(bag), "--output", str(out)])
        d = json.loads(out.read_text(encoding="utf-8"))
        assert d["dq_rules"][0]["rule_number"] == "R001"
        assert d["dq_rules"][0]["target_schema"] == "dws"
        # DQ provenance 自动采集（dq_file 同构红利）
        dq_prov = [p for p in d["provenance"] if p["type"] == "dq"]
        assert len(dq_prov) == 1 and dq_prov[0]["path"].endswith("dq.xlsx")

    def test_facts_dir_empty_dq_excel_fails_loud(self, tmp_path, capsys):
        """目录分拣出 DQ 形态但 excel 无规则行 → fail loud（错误清单退回）。"""
        bag = tmp_path / "bag"
        bag.mkdir()
        shutil.copy(ET_XLSX, bag / "et.xlsx")
        _write_dq_excel(bag / "dq.xlsx", [])
        out = tmp_path / "o.json"
        with pytest.raises(SystemExit):
            _run_facts(["--dir", str(bag), "--output", str(out)])
        assert "无有效规则行" in capsys.readouterr().err
        assert not out.exists()


class TestBrokenDimension:
    """dimension 声明容错（真实语料：平台导出写了数据但 dimension=A1 未更新）。

    Excel 人眼正常，read_only 流式读取按声明拿到空——reset_dimensions() 修复。
    构造方式：正常写 excel 后，改写 zip 内 sheet XML 的 dimension 声明为 A1。
    """

    @staticmethod
    def _break_dimension(path):
        import io
        import re as _re
        import zipfile
        src = path.read_bytes()
        buf = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                    data = _re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', data)
                zout.writestr(item, data)
        path.write_bytes(buf.getvalue())
        return path

    def test_dq_broken_dimension_still_reads(self, tmp_path):
        xlsx = tmp_path / "dq_broken.xlsx"
        _write_dq_excel(xlsx, [
            ["R001", "1", "新增", "空值检查", "主键非空", "空值检查",
             "u", "dws", "dwb_cust_asset_d", "T1", "SELECT 1", "严重"],
        ])
        self._break_dimension(xlsx)
        rules = read_dq_excel(str(xlsx))
        assert rules is not None and rules[0]["rule_number"] == "R001"

    def test_lts_broken_dimension_still_reads(self, tmp_path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "tasks"
        ws.append(["项目名称", "任务组名称", "任务名称", "任务类型", "开始时间",
                   "结束时间", "调度周期", "依赖上一周期", "日历数据", "责任人", "是否一天多调"])
        ws.append(["P", "G", "T_X_F", "周期任务", None, None, "", None, None, "u", None])
        wp = wb.create_sheet("taskParams")
        wp.append(["项目名称", "任务组名称", "任务名称", "参数名称", "参数值"])
        wp.append(["P", "G", "T_X_F", "V_GROUP_CODE", "RG1"])
        xlsx = tmp_path / "lts_broken.xlsx"
        wb.save(str(xlsx))
        self._break_dimension(xlsx)
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx), "RG1")
        assert sched is not None
        assert sched["f_tasks"][0]["task_name"] == "T_X_F"
        assert sched["f_tasks"][0]["group_code"] == "RG1"
