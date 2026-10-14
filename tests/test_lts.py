"""LTS 调度任务发现测试。

覆盖：F+I 双任务匹配 / 命名推导 / 不匹配返回 None /
损坏 yml 不崩溃 / knowledge 注入 / 干扰项不误匹配。

运行:
    pytest tests/test_lts.py -v
"""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ANALYZER_REF = PROJECT_ROOT / "dws-pipeline-analyzer" / "references"
FIXTURES = PROJECT_ROOT / "tests" / "fixtures" / "analyzer"
sys.path.insert(0, str(ANALYZER_REF))
sys.path.insert(0, str(FIXTURES))

from _build_repo import build_mock_repo
from analyzer import _discover_lts_from_repo, _parse_lts_task


@pytest.fixture
def mock_repo(tmp_path):
    """构造含真实 LTS 结构的 mock 代码仓。"""
    return build_mock_repo(tmp_path / "repo")


# ════════════════════════════════════════════════════════════════════════
# 1. 发现函数
# ════════════════════════════════════════════════════════════════════════

class TestDiscoverLts:
    """LTS 调度任务发现。"""

    def test_find_f_task_by_group_code(self, mock_repo):
        """按 V_GROUP_CODE 找到 F 任务"""
        group_dir = mock_repo["group_dir"]
        result = _discover_lts_from_repo(group_dir, "GR_TRADE_ORDER")
        assert result is not None
        assert len(result["f_tasks"]) >= 1
        assert result["f_tasks"][0]["task_name"] == "TASK_TRADE_ORDER_F"
        assert result["f_tasks"][0]["group_code"] == "GR_TRADE_ORDER"

    def test_find_i_task_by_naming(self, mock_repo):
        """从 F 任务名推导 I 任务（_F → _I）"""
        group_dir = mock_repo["group_dir"]
        result = _discover_lts_from_repo(group_dir, "GR_TRADE_ORDER")
        assert len(result["i_tasks"]) >= 1
        i_names = [t["task_name"] for t in result["i_tasks"]]
        assert "TASK_TRADE_ORDER_I" in i_names

    def test_i_task_depends_on_f(self, mock_repo):
        """I 任务的 job 依赖指向 F"""
        group_dir = mock_repo["group_dir"]
        result = _discover_lts_from_repo(group_dir, "GR_TRADE_ORDER")
        i_tasks = result["i_tasks"]
        assert len(i_tasks) >= 1
        i_jobs = i_tasks[0]["jobs"]
        # I 任务的第二个 job 是 tskdep 依赖
        dep_jobs = [j for j in i_jobs if j["type"] == "tskdep"]
        assert len(dep_jobs) >= 1

    def test_no_match_returns_none(self, mock_repo):
        """不匹配的 rule_group_code 返回 None"""
        group_dir = mock_repo["group_dir"]
        result = _discover_lts_from_repo(group_dir, "GR_NOT_EXIST")
        assert result is None

    def test_empty_group_code_returns_none(self, mock_repo):
        """空 rule_group_code 返回 None"""
        group_dir = mock_repo["group_dir"]
        assert _discover_lts_from_repo(group_dir, "") is None

    def test_outside_repo_returns_none(self, tmp_path):
        """不在代码仓里返回 None"""
        random_dir = tmp_path / "not_a_repo"
        random_dir.mkdir()
        result = _discover_lts_from_repo(random_dir, "GR_TRADE_ORDER")
        assert result is None

    def test_disturb_task_not_matched(self, mock_repo):
        """干扰任务（GROUP_OTHER）不会被误匹配"""
        group_dir = mock_repo["group_dir"]
        # GR_OTHER_NOT_MATCHED 存在于干扰 yml 里，但不该匹配 GR_TRADE_ORDER
        result = _discover_lts_from_repo(group_dir, "GR_TRADE_ORDER")
        f_names = [t["task_name"] for t in result["f_tasks"]]
        assert "TASK_TRADE_ORDER_F" in f_names
        assert "TASK_OTHER_F" not in f_names


# ════════════════════════════════════════════════════════════════════════
# 2. yml 解析
# ════════════════════════════════════════════════════════════════════════

class TestParseLtsTask:
    """单个 LTS yml 解析。"""

    def test_parse_extracts_key_fields(self, mock_repo):
        """解析出关键字段"""
        lts_file = mock_repo["repo_root"] / "LTS" / "BCNB_DAILY" / "GROUP_TRADE" / "TASK_TRADE_ORDER_F.yml"
        task = _parse_lts_task(lts_file)
        assert task is not None
        assert task["task_name"] == "TASK_TRADE_ORDER_F"
        assert task["task_type"] == "周期任务"
        assert task["schedule_cron"] == "0 30 2 * * ?"
        assert task["owner"] == "test_user"
        assert task["group_code"] == "GR_TRADE_ORDER"
        assert len(task["jobs"]) == 2
        assert len(task["params"]) == 2

    def test_parse_corrupted_yml_returns_none(self, tmp_path):
        """损坏的 yml 返回 None 不崩溃"""
        bad_file = tmp_path / "bad.yml"
        bad_file.write_text("{{{{not valid yaml", encoding="utf-8")
        assert _parse_lts_task(bad_file) is None

    def test_parse_empty_yml_returns_none(self, tmp_path):
        """空 yml 返回 None"""
        empty_file = tmp_path / "empty.yml"
        empty_file.write_text("", encoding="utf-8")
        assert _parse_lts_task(empty_file) is None


# ════════════════════════════════════════════════════════════════════════
# 3. knowledge 注入（端到端）
# ════════════════════════════════════════════════════════════════════════

class TestLtsInAnalysis:
    """LTS 注入到 knowledge（端到端集成测试）。"""

    def test_schedule_injected_into_knowledge(self, mock_repo, tmp_path):
        """分析 yml 目录后 knowledge 含 schedule"""
        import json
        from analyzer import main
        out_dir = tmp_path / "output"

        sys.argv = [
            "analyzer", "--input", str(mock_repo["group_dir"]),
            "--output", str(out_dir),
        ]
        main()

        kj_path = out_dir / mock_repo["group_dir"].name / "knowledge_draft.json"
        knowledge = json.loads(kj_path.read_text(encoding="utf-8"))

        assert "schedule" in knowledge.get("meta", {}), "knowledge.meta 应含 schedule"
        sched = knowledge["meta"]["schedule"]
        assert sched["f_tasks"][0]["task_name"] == "TASK_TRADE_ORDER_F"
        assert any(t["task_name"] == "TASK_TRADE_ORDER_I" for t in sched["i_tasks"])

    def test_no_schedule_when_xlsx_mode(self, tmp_path):
        """xlsx 模式不发现 LTS（只有 yml 模式才发现）"""
        import json
        from _build_xlsx import build_xlsx
        from analyzer import main

        xlsx_path = tmp_path / "test.xlsx"
        build_xlsx(str(xlsx_path), rules=[
            {"rule_code": "R001", "rule_type": 1, "exec_sequence": 1,
             "target_schema": "dws", "target_table": "tmp_test",
             "delete_mode": "1", "query_sql": "SELECT 1 AS id FROM dual",
             "rule_name": "test", "rule_group_code": "GR_TEST",
             "rule_group_en": "DWB_TEST_F"},
        ])
        out_dir = tmp_path / "output"
        sys.argv = ["analyzer", "--input", str(xlsx_path), "--output", str(out_dir)]
        main()

        kj_path = out_dir / "DWB_TEST_F" / "knowledge_draft.json"
        knowledge = json.loads(kj_path.read_text(encoding="utf-8"))
        assert "schedule" not in knowledge.get("meta", {}), "xlsx 模式不该有 schedule"


class TestReadLtsExcel:
    """read_lts_excel：LTS 平台导出 excel（三 sheet）→ 与 yml 发现同构（v1.8 波 1b）。

    excel 是 yml 的原始形态（yml 由平台从该 excel 转出），loader 契约：
    产出 dict 结构与 _discover_lts_from_repo 完全同构，投影层无感源差异。
    """

    @staticmethod
    def _write_lts_excel(path, tasks_rows, jobs_rows=None, params_rows=None):
        """构造 LTS 平台导出格式 excel。*_rows 不含表头。"""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "tasks"
        ws.append(["项目名称", "任务组名称", "任务名称", "任务类型", "开始时间",
                   "结束时间", "调度周期", "依赖上一周期", "日历数据", "责任人",
                   "是否一天多调"])
        for r in tasks_rows:
            ws.append(r)
        if jobs_rows is not None:
            wj = wb.create_sheet("jobs")
            wj.append(["项目名称", "任务组名称", "任务名称", "job名称", "job类型",
                       "job的父节点名称", "job超时时间", "job重试次数", "job重试间隔",
                       "job异常处理方式"])
            for r in jobs_rows:
                wj.append(r)
        if params_rows is not None:
            wp = wb.create_sheet("taskParams")
            wp.append(["项目名称", "任务组名称", "任务名称", "参数名称", "参数值"])
            for r in params_rows:
                wp.append(r)
        wb.save(str(path))

    def test_full_three_sheets_isomorphic(self, tmp_path):
        """三 sheet 完整读取，task dict 字段与 yml 解析产物同构。"""
        import datetime
        xlsx = tmp_path / "schedule.xlsx"
        self._write_lts_excel(
            xlsx,
            tasks_rows=[
                ["PROJ", "GROUP_A", "TASK_X_F", "周期任务", datetime.datetime(2026, 3, 20, 10, 31, 25),
                 None, "0 26 0 * * ?", None, None, "x00123456", None],
            ],
            jobs_rows=[
                ["PROJ", "GROUP_A", "TASK_X_F", "Pjob_x", "url", "start",
                 "10", "3", "60", "终止"],
            ],
            params_rows=[
                ["PROJ", "GROUP_A", "TASK_X_F", "V_GROUP_CODE", "RG001"],
                ["PROJ", "GROUP_A", "TASK_X_F", "V_DT", "20260320"],
            ],
        )
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx), "RG001")
        assert sched is not None
        f = sched["f_tasks"][0]
        # 与 _parse_lts_task 产物字段完全同构（loader 接口契约）
        assert set(f.keys()) == {
            "task_name", "task_type", "schedule_cron", "owner", "start_time",
            "depends_prev_cycle", "multi_day", "task_group", "project",
            "group_code", "jobs", "params", "lts_file",
        }
        assert f["task_name"] == "TASK_X_F"
        assert f["task_type"] == "周期任务"
        assert f["schedule_cron"] == "0 26 0 * * ?"
        assert f["owner"] == "x00123456"
        assert f["task_group"] == "GROUP_A"
        assert f["project"] == "PROJ"
        assert f["group_code"] == "RG001"  # taskParams 的 V_GROUP_CODE
        assert f["start_time"] == "2026-03-20 10:31:25"  # datetime → str
        assert f["jobs"] == [{"name": "Pjob_x", "type": "url", "parent": "start",
                              "retry_count": "3", "retry_interval": "60",
                              "timeout": "10", "error_handler": "终止"}]
        assert {"name": "V_DT", "value": "20260320"} in f["params"]
        assert f["lts_file"] == str(xlsx)

    def test_f_to_i_link_by_name(self, tmp_path):
        """I 任务无 V_GROUP_CODE（同 yml 场景）→ 按 _F→_I 命名关联进 i_tasks。"""
        xlsx = tmp_path / "schedule.xlsx"
        self._write_lts_excel(
            xlsx,
            tasks_rows=[
                ["PROJ", "G", "TASK_X_F", "周期任务", None, None, "", None, None, "u", None],
                ["PROJ", "G", "TASK_X_I", "周期任务", None, None, "", None, None, "u", None],
            ],
            params_rows=[
                ["PROJ", "G", "TASK_X_F", "V_GROUP_CODE", "RG001"],
            ],
        )
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx), "RG001")
        assert [t["task_name"] for t in sched["f_tasks"]] == ["TASK_X_F"]
        assert [t["task_name"] for t in sched["i_tasks"]] == ["TASK_X_I"]
        # I 任务虽不在匹配集（无 V_GROUP_CODE），仍进 all_tasks
        assert {t["task_name"] for t in sched["all_tasks"]} == {"TASK_X_F", "TASK_X_I"}

    def test_group_code_filter_multiple_groups(self, tmp_path):
        """excel 混多规则组任务时，按 V_GROUP_CODE 精确过滤。"""
        xlsx = tmp_path / "schedule.xlsx"
        self._write_lts_excel(
            xlsx,
            tasks_rows=[
                ["PROJ", "G1", "TASK_A_F", "周期任务", None, None, "", None, None, "u", None],
                ["PROJ", "G2", "TASK_B_F", "周期任务", None, None, "", None, None, "u", None],
            ],
            params_rows=[
                ["PROJ", "G1", "TASK_A_F", "V_GROUP_CODE", "RG_A"],
                ["PROJ", "G2", "TASK_B_F", "V_GROUP_CODE", "RG_B"],
            ],
        )
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx), "RG_B")
        assert [t["task_name"] for t in sched["all_tasks"]] == ["TASK_B_F"]

    def test_no_group_code_match_fallback_collect_all(self, tmp_path):
        """V_GROUP_CODE 列空值/不匹配 → 回退全收（excel 是显式输入，信任范围）。"""
        xlsx = tmp_path / "schedule.xlsx"
        self._write_lts_excel(
            xlsx,
            tasks_rows=[
                ["PROJ", "G", "TASK_X_F", "周期任务", None, None, "", None, None, "u", None],
            ],
            params_rows=[
                ["PROJ", "G", "TASK_X_F", "V_GROUP_CODE", None],  # 真实导出常见空值
            ],
        )
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx), "RG001")
        assert sched is not None
        assert len(sched["all_tasks"]) == 1
        assert sched["all_tasks"][0]["group_code"] == ""

    def test_missing_optional_sheets_explicit_empty(self, tmp_path):
        """jobs/taskParams sheet 缺失 → jobs/params 显式空列表。"""
        xlsx = tmp_path / "schedule.xlsx"
        self._write_lts_excel(
            xlsx,
            tasks_rows=[["PROJ", "G", "TASK_X_F", "周期任务", None, None, "", None, None, "u", None]],
            jobs_rows=None, params_rows=None,
        )
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx))
        t = sched["all_tasks"][0]
        assert t["jobs"] == []
        assert t["params"] == []
        assert t["group_code"] == ""

    def test_other_task_classification(self, tmp_path):
        """非 _F/_I 后缀任务归 other_tasks。"""
        xlsx = tmp_path / "schedule.xlsx"
        self._write_lts_excel(
            xlsx,
            tasks_rows=[
                ["PROJ", "G", "TASK_ODS_X", "周期任务", None, None, "", None, None, "u", None],
            ],
        )
        from inputs import read_lts_excel
        sched = read_lts_excel(str(xlsx))
        assert [t["task_name"] for t in sched["other_tasks"]] == ["TASK_ODS_X"]
        assert sched["f_tasks"] == [] and sched["i_tasks"] == []

    def test_invalid_inputs_return_none(self, tmp_path):
        from inputs import read_lts_excel
        assert read_lts_excel(tmp_path / "not_exist.xlsx") is None
        # 无 tasks sheet
        import openpyxl
        xlsx = tmp_path / "bad.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "other"
        wb.save(str(xlsx))
        assert read_lts_excel(str(xlsx)) is None
